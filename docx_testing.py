from docx import Document
from docx import table
from pymongo import MongoClient
from pprint import pprint
from collections import defaultdict
import requests
import os
from pathlib import Path
import time
directory = os.getcwd()
database = "http://localhost:3000/v1/skills"

for filename in os.listdir(directory):
    if filename.endswith('.docx'):

        document = Document(Path(filename).absolute())

        table = document.tables[0]

        # GET HEADER INFORMATION (Title, task, task #, etc)
        titleLocation = (0, 0)
        evaluatorInstructionsLocation = (1, 0)
        taskLocation = (2, 0)
        performanceOutcomeLocation = (3, 0)
        candidateDirectivesLocation = (4, 0)

        taskList = []

        currentPost = {
            'title': table.cell(*titleLocation).text,
            'evaluator_instructions': table.cell(*evaluatorInstructionsLocation).text,
            'candidate_driectives': table.cell(*candidateDirectivesLocation).text,
        }
        docInfo = [table.cell(*titleLocation).text, table.cell(*evaluatorInstructionsLocation).text,
            table.cell(*taskLocation).text, table.cell(*performanceOutcomeLocation).text, table.cell(*candidateDirectivesLocation).text]

        for row in range(7, len(table.rows) - 1): # get tasks remove last row
            print(table.cell(row, 1).text)    
            taskList.append(table.cell(row, 1).text)

        currentPost['tasks'] = taskList#add tasks in list format to dictionary with the key 'tasks'

        print(table.cell(*taskLocation).text)

        r = requests.post(url = database, json = currentPost)

        
    else:
        continue
